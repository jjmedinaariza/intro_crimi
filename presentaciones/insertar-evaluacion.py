"""
Inserta la diapositiva de evaluación en la posición 7 del deck YA MODIFICADO.
No regenera nada: abre el .pptx existente y le añade una diapositiva.
Datos: Proyecto Docente 2180003, curso 2026-27.
"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC = "/sessions/clever-modest-cray/mnt/intro_crimi/presentaciones/Cap01_Que-es-la-criminologia.pptx"
POS = 7  # la nueva diapositiva pasa a ocupar esta posición (1-based)

# Tokens "B claro"
INK = RGBColor(0x14, 0x17, 0x1A)
MUTED = RGBColor(0x5F, 0x66, 0x6D)
ACCENT = RGBColor(0xA6, 0x6A, 0x00)   # ocre: acento sobre blanco
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

M = 0.95
CW = 13.333 - 2 * M


def caja(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def escribir(tf, lineas, size, color, bold=False, spacing=None):
    """lineas: lista de strings, una por párrafo."""
    for i, texto in enumerate(lineas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = texto
        f = r.font
        f.name = "Arial"
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = color
        if spacing:
            p.line_spacing = Pt(spacing)


prs = Presentation(SRC)
layout = prs.slides[6].slide_layout          # el mismo layout DEFAULT del resto
sl = prs.slides.add_slide(layout)

# Fondo blanco explícito (el resto del deck lo lleva)
fill = sl.background.fill
fill.solid()
fill.fore_color.rgb = WHITE

# --- titular ---------------------------------------------------------------
tf = caja(sl, M, 1.05, CW, 0.8)
escribir(tf, ["Cómo se evalúa"], 38, INK, bold=True)

# --- dos columnas ----------------------------------------------------------
cols = [
    ("30 %  ·  Evaluación continua",
     ["Al menos 6 actividades.",
      "Puntúan de 1 a 10 y pesan igual.",
      "Dos son controles de lectura."]),
    ("70 %  ·  Examen final",
     ["No es tipo test.",
      "Preguntas para desarrollar.",
      "La nota final es la media ponderada."]),
]
for i, (rotulo, lineas) in enumerate(cols):
    x = M + i * 5.95
    tf = caja(sl, x, 2.3, 5.4, 0.55)
    escribir(tf, [rotulo], 23, ACCENT, bold=True)
    tf = caja(sl, x, 2.98, 5.4, 1.6)
    escribir(tf, lineas, 17, INK, spacing=29)

# --- la vía alternativa ----------------------------------------------------
tf = caja(sl, M, 4.78, 11.0, 0.5)
escribir(tf, ["¿No haces la continua? Puedes jugártelo todo al examen final — "
              "avisando antes al coordinador."], 16, MUTED)

# --- notas del ponente -----------------------------------------------------
# El notesMaster de pptxgenjs no trae placeholders, así que la diapositiva de
# notas nueva sale vacía: se clona la estructura de una existente y se vacía.
src_ns = prs.slides[6].notes_slide
dst_ns = sl.notes_slide
dst_tree = dst_ns.shapes._spTree
src_tree = src_ns.shapes._spTree
for child in list(dst_tree):
    if child.tag.endswith("}sp") or child.tag.endswith("}pic"):
        dst_tree.remove(child)
for child in src_tree:
    if child.tag.endswith("}sp") or child.tag.endswith("}pic"):
        dst_tree.append(copy.deepcopy(child))

dst_ns.notes_text_frame.text = (
    "ES LO QUE MÁS LES IMPORTA. No pasar rápido: dejar que copien esta y solo esta.\n\n"
    "LA CONTINUA (30 %). Serán al menos 6 actividades, todas puntuadas del 1 al 10 y todas con "
    "el mismo peso entre sí. Al menos 2 están diseñadas como controles de lectura — decirlo "
    "explícitamente, porque es la palanca para que lean el manual.\n\n"
    "EL EXAMEN FINAL (70 %). Insistir en que NO es tipo test: incluye preguntas que hay que "
    "desarrollar. Suele ser la pregunta número uno y conviene adelantarse a ella.\n\n"
    "LA NOTA. Media ponderada entre los dos componentes, sin más.\n\n"
    "LA VÍA ALTERNATIVA. Quien no quiera hacer la continua puede jugárselo todo al examen final, "
    "pero con dos condiciones que hay que subrayar: (1) tiene que informar previamente al "
    "coordinador de la asignatura, y (2) hará un examen final DISTINTO Y MÁS LARGO que quienes sí "
    "hicieron la continua. No es la misma prueba.\n\n"
    "Enlazar con la diapositiva de 'estas diapositivas no son los apuntes': si hay controles de "
    "lectura, el sitio donde se estudia es el manual, no la pantalla.\n\n"
    "Fuente: Proyecto Docente 2180003 «Introducción a la Criminología», curso 2026-27, "
    "última modificación 29/06/2026."
)

# --- mover a la posición POS ----------------------------------------------
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
sldIdLst.remove(ids[-1])
sldIdLst.insert(POS - 1, ids[-1])

prs.save(SRC)

# --- parche de empaquetado -------------------------------------------------
# Al reguardar, python-pptx descarta el <Default Extension="jpg"> que había
# escrito pptxgenjs (con el ContentType no estándar "image/jpg"). Sin él, las
# dos imágenes .jpg del deck —el retrato de Sutherland y la foto de BLM— se
# quedan sin tipo declarado y PowerPoint puede no abrirlas.
import shutil
import zipfile

CT = "[Content_Types].xml"
tmp = SRC + ".fix"
with zipfile.ZipFile(SRC) as zin:
    datos = {n: zin.read(n) for n in zin.namelist()}

ct = datos[CT].decode("utf-8")
if 'Extension="jpg"' not in ct:
    ct = ct.replace(
        '<Default Extension="png"',
        '<Default Extension="jpg" ContentType="image/jpeg"/>'
        '<Default Extension="png"',
        1,
    )
    datos[CT] = ct.encode("utf-8")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, d in datos.items():
            zout.writestr(n, d)
    shutil.move(tmp, SRC)
    print("· [Content_Types].xml: repuesto el Default de jpg")

print(f"OK · insertada en la posición {POS} · total {len(sldIdLst)} diapositivas")
